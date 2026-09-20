"""
Personal Financial Digital Twin — PDF and Excel Export Service Module

Generates professional, formatted PDF reports (via ReportLab) and multi-worksheet
Excel workbooks (via openpyxl) representing the user's current financial inputs
and application outputs using plain numeric values.

Stateless and non-duplicative — receives current frontend/application state,
formats the document in-memory, and returns raw binary bytes.
"""

import io
from datetime import datetime
from typing import Dict, Any, Optional, List

# --- ReportLab Imports ---
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

# --- OpenPyXL Imports ---
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


FRIENDLY_NAMES = {
    "ending_balance_t": "Current Month Ending Balance",
    "income_credit_t": "Income & Inflows",
    "debit_count_t": "Number of Payments",
    "spending_hh_t": "Household Spending",
    "spending_st_t": "Bank Fees & Statements",
    "spending_in_t": "Insurance Payments",
    "spending_lo_t": "Loan Repayments",
    "spending_io_t": "Interest Charges",
    "spending_other_t": "Other Outflows",
    "spending_t_minus_1": "Last Month's Spending",
    "spending_t_minus_2": "Spending 2 Months Ago",
    "spending_t": "Current Month Spending (Derived)",
    "spending_3m_mean": "3-Month Average Spending (Derived)",
    "spending_3m_std": "3-Month Spending Variation (Derived)",
}


def _format_number(val: Any) -> str:
    if val is None or not isinstance(val, (int, float)):
        return "N/A"
    return f"{val:,.2f}"


def generate_pdf_report(data: Dict[str, Any]) -> bytes:
    """
    Generates a structured PDF report in-memory using plain numeric formatting.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    # Custom Palette
    c_primary = colors.HexColor("#0F172A")    # Slate 900
    c_secondary = colors.HexColor("#334155")  # Slate 700
    c_accent = colors.HexColor("#0284C7")     # Sky 600
    c_bg_light = colors.HexColor("#F8FAFC")   # Slate 50
    c_border = colors.HexColor("#E2E8F0")     # Slate 200
    c_pos = colors.HexColor("#16A34A")        # Green 600
    c_neg = colors.HexColor("#DC2626")        # Red 600

    # Custom Typography Styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=c_primary,
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=c_secondary,
        spaceAfter=12,
    )
    h2_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=c_primary,
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "BodyDark",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=c_secondary,
    )
    meta_style = ParagraphStyle(
        "MetaText",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#64748B"),
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=c_secondary,
    )
    cell_bold = ParagraphStyle(
        "TableCellBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=c_primary,
    )

    story = []

    # 1. Header Banner
    gen_time = data.get("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    header_table_data = [
        [
            Paragraph("Personal Financial Digital Twin", title_style),
            Paragraph(f"<b>Report Generated:</b><br/>{gen_time}", meta_style),
        ],
        [
            Paragraph("Explainable Spending Forecasting & Financial Twin Report", subtitle_style),
            Paragraph("CatBoost Model Projections", meta_style),
        ],
    ]
    header_table = Table(header_table_data, colWidths=[380, 160])
    header_table.setStyle(
        TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ])
    )
    story.append(header_table)
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_accent, spaceBefore=4, spaceAfter=10))

    # 2. Key Forecast Summary Card
    prediction = data.get("prediction")
    pred_str = _format_number(prediction) if prediction is not None else "--"
    
    summary_box_data = [
        [
            Paragraph("<b>Next-Month Spending Forecast</b>", cell_bold),
            Paragraph(f"<font size=14 color='{c_accent.hexval()}'><b>{pred_str}</b></font>", cell_bold),
        ]
    ]
    summary_box = Table(summary_box_data, colWidths=[300, 240])
    summary_box.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), c_bg_light),
            ("BOX", (0, 0), (-1, -1), 1, c_border),
            ("PADDING", (0, 0), (-1, -1), 10),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ])
    )
    story.append(summary_box)
    story.append(Spacer(1, 10))

    # 3. Current Financial Inputs Section
    story.append(Paragraph("1. Current Financial Position Inputs", h2_style))
    baseline = data.get("baseline_state", {})
    
    if baseline:
        # Construct derived preview values
        hh = baseline.get("spending_hh_t", 0.0)
        st = baseline.get("spending_st_t", 0.0)
        in_val = baseline.get("spending_in_t", 0.0)
        lo = baseline.get("spending_lo_t", 0.0)
        io_val = baseline.get("spending_io_t", 0.0)
        oth = baseline.get("spending_other_t", 0.0)
        tm1 = baseline.get("spending_t_minus_1", 0.0)
        tm2 = baseline.get("spending_t_minus_2", 0.0)

        spending_t = hh + st + in_val + lo + io_val + oth
        mean_3m = (spending_t + tm1 + tm2) / 3.0
        var_3m = ((spending_t - mean_3m)**2 + (tm1 - mean_3m)**2 + (tm2 - mean_3m)**2) / 3.0
        std_3m = var_3m ** 0.5

        inputs_table_data = [
            [Paragraph("<b>Input Parameter</b>", cell_bold), Paragraph("<b>Value</b>", cell_bold), Paragraph("<b>Category / Type</b>", cell_bold)],
            [Paragraph(FRIENDLY_NAMES["ending_balance_t"], cell_style), Paragraph(_format_number(baseline.get("ending_balance_t")), cell_style), Paragraph("Primary Snapshot", cell_style)],
            [Paragraph(FRIENDLY_NAMES["income_credit_t"], cell_style), Paragraph(_format_number(baseline.get("income_credit_t")), cell_style), Paragraph("Primary Snapshot", cell_style)],
            [Paragraph(FRIENDLY_NAMES["debit_count_t"], cell_style), Paragraph(str(baseline.get("debit_count_t", 0)), cell_style), Paragraph("Payment Frequency", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_hh_t"], cell_style), Paragraph(_format_number(hh), cell_style), Paragraph("Category Outflow", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_st_t"], cell_style), Paragraph(_format_number(st), cell_style), Paragraph("Category Outflow", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_in_t"], cell_style), Paragraph(_format_number(in_val), cell_style), Paragraph("Category Outflow", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_lo_t"], cell_style), Paragraph(_format_number(lo), cell_style), Paragraph("Category Outflow", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_io_t"], cell_style), Paragraph(_format_number(io_val), cell_style), Paragraph("Category Outflow", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_other_t"], cell_style), Paragraph(_format_number(oth), cell_style), Paragraph("Category Outflow", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_t_minus_1"], cell_style), Paragraph(_format_number(tm1), cell_style), Paragraph("Recent History (t-1)", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_t_minus_2"], cell_style), Paragraph(_format_number(tm2), cell_style), Paragraph("Recent History (t-2)", cell_style)],
            [Paragraph(FRIENDLY_NAMES["spending_t"], cell_bold), Paragraph(_format_number(spending_t), cell_bold), Paragraph("Calculated Derived", cell_bold)],
            [Paragraph(FRIENDLY_NAMES["spending_3m_mean"], cell_bold), Paragraph(_format_number(mean_3m), cell_bold), Paragraph("Calculated Derived", cell_bold)],
            [Paragraph(FRIENDLY_NAMES["spending_3m_std"], cell_bold), Paragraph(_format_number(std_3m), cell_bold), Paragraph("Calculated Derived", cell_bold)],
        ]

        t_inputs = Table(inputs_table_data, colWidths=[240, 160, 140])
        t_inputs.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), c_bg_light),
                ("BOX", (0, 0), (-1, -1), 0.5, c_border),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, c_border),
                ("PADDING", (0, 0), (-1, -1), 4),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BACKGROUND", (0, 12), (-1, 14), colors.HexColor("#EFF6FF")),  # Highlight derived rows
            ])
        )
        story.append(t_inputs)

    story.append(Spacer(1, 12))

    # 4. What-If Scenario Section (Optional)
    scenario_data = data.get("scenario_data")
    if scenario_data and isinstance(scenario_data, dict):
        scen_story = []
        scen_story.append(Paragraph("2. What-If Scenario Analysis", h2_style))
        
        base_val = scenario_data.get("basePred")
        base_p = base_val if base_val is not None else scenario_data.get("baseline_prediction")

        scen_val = scenario_data.get("scenPred")
        scen_p = scen_val if scen_val is not None else scenario_data.get("scenario_prediction")

        abs_val = scenario_data.get("absDiff")
        abs_d = abs_val if abs_val is not None else scenario_data.get("absolute_difference")

        pct_val = scenario_data.get("pctDiff")
        pct_d = pct_val if pct_val is not None else scenario_data.get("percentage_difference")

        pct_str = f"{pct_d:+.2f}%" if pct_d is not None else "N/A"
        abs_str = f"{abs_d:+.2f}" if abs_d is not None else "--"

        scen_table_data = [
            [Paragraph("<b>Metric</b>", cell_bold), Paragraph("<b>Baseline Estimate</b>", cell_bold), Paragraph("<b>Scenario Estimate</b>", cell_bold), Paragraph("<b>Difference</b>", cell_bold)],
            [Paragraph("Forecast Spending", cell_style), Paragraph(_format_number(base_p), cell_style), Paragraph(_format_number(scen_p), cell_style), Paragraph(f"<b>{abs_str} ({pct_str})</b>", cell_style)],
        ]
        t_scen = Table(scen_table_data, colWidths=[150, 130, 130, 130])
        t_scen.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), c_bg_light),
                ("BOX", (0, 0), (-1, -1), 0.5, c_border),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, c_border),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ])
        )
        scen_story.append(t_scen)
        scen_story.append(Spacer(1, 10))
        story.append(KeepTogether(scen_story))

    # 5. Savings Target Plan Section (Optional)
    savings_plan = data.get("savings_plan")
    if savings_plan and isinstance(savings_plan, dict):
        sav_story = []
        sav_story.append(Paragraph("3. Savings Goal & Plan Schedule", h2_style))

        goal_name = savings_plan.get("rawName") or savings_plan.get("goalName", "Savings Goal")
        target_amt = savings_plan.get("amount", 0.0)
        timeframe = savings_plan.get("months", 1)
        req_monthly = savings_plan.get("requiredMonthly", 0.0)
        schedule = savings_plan.get("schedule", [])

        sav_meta_data = [
            [Paragraph(f"<b>Goal Name:</b> {goal_name}", cell_style), Paragraph(f"<b>Target Goal:</b> {_format_number(target_amt)}", cell_style)],
            [Paragraph(f"<b>Timeframe:</b> {timeframe} months", cell_style), Paragraph(f"<b>Required Monthly Saving:</b> {_format_number(req_monthly)}", cell_bold)],
        ]
        t_sav_meta = Table(sav_meta_data, colWidths=[270, 270])
        t_sav_meta.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), c_bg_light),
                ("BOX", (0, 0), (-1, -1), 0.5, c_border),
                ("PADDING", (0, 0), (-1, -1), 6),
            ])
        )
        sav_story.append(t_sav_meta)
        sav_story.append(Spacer(1, 6))

        if schedule:
            sched_table_data = [[Paragraph("<b>Month</b>", cell_bold), Paragraph("<b>Planned Monthly Saving</b>", cell_bold), Paragraph("<b>Remaining Balance</b>", cell_bold)]]
            for row in schedule[:12]:  # Show up to first 12 months for brevity in PDF
                m_num = row.get("month")
                p_sav = row.get("plannedSaving", 0.0)
                r_bal = row.get("remainingGoal", 0.0)
                sched_table_data.append([
                    Paragraph(f"Month {m_num}", cell_style),
                    Paragraph(_format_number(p_sav), cell_style),
                    Paragraph(_format_number(r_bal), cell_style),
                ])

            if len(schedule) > 12:
                sched_table_data.append([Paragraph(f"... ({len(schedule) - 12} additional months remaining in plan)", meta_style), Paragraph("", cell_style), Paragraph("", cell_style)])

            t_sched = Table(sched_table_data, colWidths=[140, 200, 200])
            t_sched.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), c_bg_light),
                    ("BOX", (0, 0), (-1, -1), 0.5, c_border),
                    ("INNERGRID", (0, 0), (-1, -1), 0.5, c_border),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ])
            )
            sav_story.append(t_sched)
            sav_story.append(Spacer(1, 10))

        story.append(KeepTogether(sav_story))

    # 6. TreeSHAP Explanation Section (Optional)
    shap_data = data.get("shap_data")
    if shap_data and isinstance(shap_data, dict):
        shap_story = []
        shap_story.append(Paragraph("4. TreeSHAP Model Feature Attributions", h2_style))

        base_val = shap_data.get("base_value", 0.0)
        recon_val = shap_data.get("reconstructed_prediction", 0.0)
        delta_val = shap_data.get("additivity_delta", 0.0)
        features = shap_data.get("features", [])

        shap_meta = Paragraph(
            f"<b>Dataset Expected Base Value E[f(X)]:</b> {_format_number(base_val)} &nbsp;|&nbsp; "
            f"<b>Reconstructed:</b> {_format_number(recon_val)} &nbsp;|&nbsp; "
            f"<b>Additivity Delta:</b> {delta_val:.6f}",
            meta_style
        )
        shap_story.append(shap_meta)
        shap_story.append(Spacer(1, 6))

        if features:
            shap_table_data = [[Paragraph("<b>Feature</b>", cell_bold), Paragraph("<b>Feature Value</b>", cell_bold), Paragraph("<b>SHAP Value (Contribution)</b>", cell_bold)]]
            # Sort by absolute SHAP impact
            sorted_feats = sorted(features, key=lambda f: abs(f.get("shap_value", 0.0)), reverse=True)
            for f in sorted_feats:
                fname = FRIENDLY_NAMES.get(f.get("feature"), f.get("feature"))
                fval = f.get("value")
                fval_str = str(fval) if isinstance(fval, int) else f"{fval:,.2f}" if isinstance(fval, float) else str(fval)
                sval = f.get("shap_value", 0.0)
                color_hex = c_pos.hexval() if sval >= 0 else c_neg.hexval()
                sval_str = f"<font color='{color_hex}'><b>{sval:+.4f}</b></font>"

                shap_table_data.append([
                    Paragraph(fname, cell_style),
                    Paragraph(fval_str, cell_style),
                    Paragraph(sval_str, cell_style),
                ])

            t_shap = Table(shap_table_data, colWidths=[240, 150, 150])
            t_shap.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), c_bg_light),
                    ("BOX", (0, 0), (-1, -1), 0.5, c_border),
                    ("INNERGRID", (0, 0), (-1, -1), 0.5, c_border),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ])
            )
            shap_story.append(t_shap)
            shap_story.append(Spacer(1, 10))

        story.append(KeepTogether(shap_story))

    # Footer Disclaimer
    story.append(HRFlowable(width="100%", thickness=0.5, color=c_border, spaceBefore=10, spaceAfter=8))
    story.append(
        Paragraph(
            "<b>Disclaimer:</b> Model spending forecasts and scenario analyses are statistical estimates generated by a trained CatBoost machine learning model. "
            "They represent analytical model projections and do not guarantee actual future financial outcomes.",
            meta_style,
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(data: Dict[str, Any]) -> bytes:
    """
    Generates a multi-worksheet Excel (.xlsx) workbook in-memory using plain numeric formatting.
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    gen_time = data.get("generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    # Styles & Colors
    font_family = "Arial"
    font_title = Font(name=font_family, size=14, bold=True, color="0F172A")
    font_subtitle = Font(name=font_family, size=10, italic=True, color="475569")
    font_header = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    font_bold = Font(name=font_family, size=9.5, bold=True, color="0F172A")
    font_body = Font(name=font_family, size=9.5, color="334155")
    font_pos = Font(name=font_family, size=9.5, bold=True, color="16A34A")
    font_neg = Font(name=font_family, size=9.5, bold=True, color="DC2626")

    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_subheader = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_derived = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    border_thin = Side(border_style="thin", color="CBD5E1")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    num_format_currency = '#,##0.00'
    num_format_number = '#,##0.00'
    num_format_int = '#,##0'

    # -------------------------------------------------------------
    # WORKSHEET 1: Executive Summary
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum["A1"] = "Personal Financial Digital Twin — Executive Summary"
    ws_sum["A1"].font = font_title
    ws_sum["A2"] = f"Generated: {gen_time}"
    ws_sum["A2"].font = font_subtitle

    prediction = data.get("prediction")
    baseline = data.get("baseline_state", {})

    ws_sum.append([])
    ws_sum.append(["Key Metrics Overview", "Value"])
    row_idx = 4
    for col in range(1, 3):
        cell = ws_sum.cell(row=row_idx, column=col)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_left if col == 1 else align_right

    summary_rows = [
        ("Next-Month Spending Forecast", prediction, num_format_currency),
        ("Current Month Ending Balance", baseline.get("ending_balance_t"), num_format_currency),
        ("Income & Inflows", baseline.get("income_credit_t"), num_format_currency),
        ("Number of Payments", baseline.get("debit_count_t"), num_format_int),
    ]

    for label, val, fmt in summary_rows:
        row_idx += 1
        c1 = ws_sum.cell(row=row_idx, column=1, value=label)
        c2 = ws_sum.cell(row=row_idx, column=2, value=val)
        c1.font = font_bold if label == "Next-Month Spending Forecast" else font_body
        c2.font = font_bold if label == "Next-Month Spending Forecast" else font_body
        c1.border = box_border
        c2.border = box_border
        c1.alignment = align_left
        c2.alignment = align_right
        if isinstance(val, (int, float)):
            c2.number_format = fmt

    # Optional section highlights
    row_idx += 2
    ws_sum.cell(row=row_idx, column=1, value="Module Summary").font = font_title
    row_idx += 1
    ws_sum.cell(row=row_idx, column=1, value="Module").font = font_header
    ws_sum.cell(row=row_idx, column=1).fill = fill_header
    ws_sum.cell(row=row_idx, column=2, value="Status / Highlight").font = font_header
    ws_sum.cell(row=row_idx, column=2).fill = fill_header

    scen_d = data.get("scenario_data")
    sav_d = data.get("savings_plan")
    shap_d = data.get("shap_data")

    modules = [
        ("What-If Counterfactual Analysis", "Active" if scen_d else "Not Executed"),
        ("Savings Goal Plan Schedule", f"Goal: {sav_d.get('rawName') or sav_d.get('goalName')}" if sav_d else "Not Executed"),
        ("TreeSHAP Feature Attributions", "Generated" if shap_d else "Not Executed"),
    ]

    for m_name, m_status in modules:
        row_idx += 1
        c1 = ws_sum.cell(row=row_idx, column=1, value=m_name)
        c2 = ws_sum.cell(row=row_idx, column=2, value=m_status)
        c1.font = font_body
        c2.font = font_body
        c1.border = box_border
        c2.border = box_border

    # -------------------------------------------------------------
    # WORKSHEET 2: Financial Inputs
    # -------------------------------------------------------------
    ws_inp = wb.create_sheet(title="Financial Inputs")
    ws_inp.views.sheetView[0].showGridLines = True

    ws_inp["A1"] = "Financial State Inputs & Derived Feature Schema"
    ws_inp["A1"].font = font_title

    ws_inp.append([])
    ws_inp.append(["Feature Name", "Friendly Label", "Value", "Type / Category"])
    header_row = 3
    for col in range(1, 5):
        cell = ws_inp.cell(row=header_row, column=col)
        cell.font = font_header
        cell.fill = fill_header

    # Primary fields
    fields_order = [
        ("ending_balance_t", "Primary Snapshot", num_format_currency),
        ("income_credit_t", "Primary Snapshot", num_format_currency),
        ("debit_count_t", "Payment Frequency", num_format_int),
        ("spending_hh_t", "Category Outflow", num_format_currency),
        ("spending_st_t", "Category Outflow", num_format_currency),
        ("spending_in_t", "Category Outflow", num_format_currency),
        ("spending_lo_t", "Category Outflow", num_format_currency),
        ("spending_io_t", "Category Outflow", num_format_currency),
        ("spending_other_t", "Category Outflow", num_format_currency),
        ("spending_t_minus_1", "Recent History (t-1)", num_format_currency),
        ("spending_t_minus_2", "Recent History (t-2)", num_format_currency),
    ]

    curr_row = 3
    for fkey, fcat, fmt in fields_order:
        curr_row += 1
        val = baseline.get(fkey, 0.0)
        c1 = ws_inp.cell(row=curr_row, column=1, value=fkey)
        c2 = ws_inp.cell(row=curr_row, column=2, value=FRIENDLY_NAMES.get(fkey, fkey))
        c3 = ws_inp.cell(row=curr_row, column=3, value=val)
        c4 = ws_inp.cell(row=curr_row, column=4, value=fcat)

        for cell in (c1, c2, c3, c4):
            cell.font = font_body
            cell.border = box_border
        c3.alignment = align_right
        if isinstance(val, (int, float)):
            c3.number_format = fmt

    # Calculate derived rows
    hh = baseline.get("spending_hh_t", 0.0)
    st = baseline.get("spending_st_t", 0.0)
    in_val = baseline.get("spending_in_t", 0.0)
    lo = baseline.get("spending_lo_t", 0.0)
    io_val = baseline.get("spending_io_t", 0.0)
    oth = baseline.get("spending_other_t", 0.0)
    tm1 = baseline.get("spending_t_minus_1", 0.0)
    tm2 = baseline.get("spending_t_minus_2", 0.0)

    s_t = hh + st + in_val + lo + io_val + oth
    mean_3m = (s_t + tm1 + tm2) / 3.0
    var_3m = ((s_t - mean_3m)**2 + (tm1 - mean_3m)**2 + (tm2 - mean_3m)**2) / 3.0
    std_3m = var_3m ** 0.5

    derived_fields = [
        ("spending_t", FRIENDLY_NAMES["spending_t"], s_t, num_format_currency),
        ("spending_3m_mean", FRIENDLY_NAMES["spending_3m_mean"], mean_3m, num_format_currency),
        ("spending_3m_std", FRIENDLY_NAMES["spending_3m_std"], std_3m, num_format_number),
    ]

    for fkey, flabel, val, fmt in derived_fields:
        curr_row += 1
        c1 = ws_inp.cell(row=curr_row, column=1, value=fkey)
        c2 = ws_inp.cell(row=curr_row, column=2, value=flabel)
        c3 = ws_inp.cell(row=curr_row, column=3, value=val)
        c4 = ws_inp.cell(row=curr_row, column=4, value="Calculated Derived")

        for cell in (c1, c2, c3, c4):
            cell.font = font_bold
            cell.fill = fill_derived
            cell.border = box_border
        c3.alignment = align_right
        if isinstance(val, (int, float)):
            c3.number_format = fmt

    # -------------------------------------------------------------
    # WORKSHEET 3: Forecast & Scenario
    # -------------------------------------------------------------
    if scen_d:
        ws_scen = wb.create_sheet(title="Forecast & Scenario")
        ws_scen.views.sheetView[0].showGridLines = True

        ws_scen["A1"] = "What-If Counterfactual Scenario Analysis"
        ws_scen["A1"].font = font_title

        ws_scen.append([])
        ws_scen.append(["Metric / Feature", "Baseline Value", "Scenario Value", "Absolute Difference", "Percentage Change"])
        h_row = 3
        for col in range(1, 6):
            cell = ws_scen.cell(row=h_row, column=col)
            cell.font = font_header
            cell.fill = fill_header

        base_val = scen_d.get("basePred")
        b_p = base_val if base_val is not None else scen_d.get("baseline_prediction", 0.0)

        scen_val = scen_d.get("scenPred")
        s_p = scen_val if scen_val is not None else scen_d.get("scenario_prediction", 0.0)

        abs_val = scen_d.get("absDiff")
        a_d = abs_val if abs_val is not None else scen_d.get("absolute_difference", 0.0)

        pct_val = scen_d.get("pctDiff")
        p_d = pct_val if pct_val is not None else scen_d.get("percentage_difference")

        ws_scen.append(["Next-Month Spending Forecast", b_p, s_p, a_d, p_d / 100.0 if p_d is not None else "N/A"])
        r_idx = 4
        c1 = ws_scen.cell(row=r_idx, column=1)
        c2 = ws_scen.cell(row=r_idx, column=2)
        c3 = ws_scen.cell(row=r_idx, column=3)
        c4 = ws_scen.cell(row=r_idx, column=4)
        c5 = ws_scen.cell(row=r_idx, column=5)

        for c in (c1, c2, c3, c4, c5):
            c.font = font_bold
            c.border = box_border
        c2.number_format = num_format_currency
        c3.number_format = num_format_currency
        c4.number_format = num_format_currency
        if isinstance(p_d, (int, float)):
            c5.number_format = "0.00%"

    # -------------------------------------------------------------
    # WORKSHEET 4: Savings Goal Schedule
    # -------------------------------------------------------------
    if sav_d:
        ws_sav = wb.create_sheet(title="Savings Goal Schedule")
        ws_sav.views.sheetView[0].showGridLines = True

        ws_sav["A1"] = f"Savings Target Plan: {sav_d.get('rawName') or sav_d.get('goalName', 'Goal')}"
        ws_sav["A1"].font = font_title

        ws_sav.append([])
        ws_sav.append(["Parameter", "Value"])
        ws_sav.append(["Target Goal Amount", sav_d.get("amount", 0.0)])
        ws_sav.append(["Timeframe (Months)", sav_d.get("months", 1)])
        ws_sav.append(["Required Monthly Saving", sav_d.get("requiredMonthly", 0.0)])

        for r in range(3, 6):
            c1 = ws_sav.cell(row=r, column=1)
            c2 = ws_sav.cell(row=r, column=2)
            c1.font = font_bold
            c2.font = font_bold
            c1.border = box_border
            c2.border = box_border
            if r in (3, 5):
                c2.number_format = num_format_currency

        ws_sav.append([])
        ws_sav.append(["Month #", "Planned Monthly Saving", "Remaining Goal Balance"])
        s_h_row = 7
        for col in range(1, 4):
            cell = ws_sav.cell(row=s_h_row, column=col)
            cell.font = font_header
            cell.fill = fill_header

        schedule = sav_d.get("schedule", [])
        for row_data in schedule:
            s_r_idx = ws_sav.max_row + 1
            c1 = ws_sav.cell(row=s_r_idx, column=1, value=f"Month {row_data.get('month')}")
            c2 = ws_sav.cell(row=s_r_idx, column=2, value=row_data.get("plannedSaving", 0.0))
            c3 = ws_sav.cell(row=s_r_idx, column=3, value=row_data.get("remainingGoal", 0.0))

            for cell in (c1, c2, c3):
                cell.font = font_body
                cell.border = box_border
            c2.number_format = num_format_currency
            c3.number_format = num_format_currency

    # -------------------------------------------------------------
    # WORKSHEET 5: SHAP Attributions
    # -------------------------------------------------------------
    if shap_d:
        ws_shap = wb.create_sheet(title="SHAP Attributions")
        ws_shap.views.sheetView[0].showGridLines = True

        ws_shap["A1"] = "TreeSHAP Local Feature Attributions"
        ws_shap["A1"].font = font_title

        ws_shap["A2"] = f"Expected Base Value E[f(X)]: {shap_d.get('base_value', 0.0):,.2f} | Reconstructed: {shap_d.get('reconstructed_prediction', 0.0):,.2f} | Additivity Delta: {shap_d.get('additivity_delta', 0.0):.6f}"
        ws_shap["A2"].font = font_subtitle

        ws_shap.append([])
        ws_shap.append(["Feature Identifier", "Friendly Label", "Feature Value", "SHAP Value (Contribution)"])
        sh_row = 4
        for col in range(1, 5):
            cell = ws_shap.cell(row=sh_row, column=col)
            cell.font = font_header
            cell.fill = fill_header

        features = shap_d.get("features", [])
        sorted_feats = sorted(features, key=lambda f: abs(f.get("shap_value", 0.0)), reverse=True)

        for f in sorted_feats:
            fkey = f.get("feature")
            flabel = FRIENDLY_NAMES.get(fkey, fkey)
            fval = f.get("value")
            sval = f.get("shap_value", 0.0)

            row_n = ws_shap.max_row + 1
            c1 = ws_shap.cell(row=row_n, column=1, value=fkey)
            c2 = ws_shap.cell(row=row_n, column=2, value=flabel)
            c3 = ws_shap.cell(row=row_n, column=3, value=fval)
            c4 = ws_shap.cell(row=row_n, column=4, value=sval)

            c1.font = font_body
            c2.font = font_body
            c3.font = font_body
            c4.font = font_pos if sval >= 0 else font_neg

            for c in (c1, c2, c3, c4):
                c.border = box_border
            c3.alignment = align_right
            c4.alignment = align_right
            if isinstance(fval, (int, float)):
                c3.number_format = num_format_number
            c4.number_format = "+#,##0.0000;-#,##0.0000;0.0000"

    # Auto-fit column widths across all worksheets
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    if cell.row in (1, 2):
                        continue
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
